from typing import List
from pathlib import Path
from langchain_core.documents import Document
import csv

class DocumentLoader:
    #parent class that all the specific loaders inherit from
    
    def load(self, file_paths: List[str]) -> List[Document]:
        raise NotImplementedError("Each loader needs its own load method")

class PDFLoader(DocumentLoader):
    #handles PDF files, extracts all pages into one document
    
    def load(self, file_paths: List[str]) -> List[Document]:
        docs = []

        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("You need pypdf installed: pip install pypdf")

        for p in file_paths:
            if not p.lower().endswith(".pdf"):
                continue

            reader = PdfReader(p)
            pages = []

            #pull text from each page
            for page in reader.pages:
                text = page.extract_text() or ""
                if text:
                    pages.append(text)

            full_text = "\n".join(pages).strip()
            if not full_text:
                continue

            docs.append(
                Document(
                    page_content=full_text,
                    metadata={
                        "source": Path(p).name,
                        "file_path": p,
                        "type": "pdf"
                    }
                )
            )
        return docs

class CSVLoader(DocumentLoader):
    #reads CSV files row by row, each row becomes a separate document
    
    def load(self, file_paths: List[str]) -> List[Document]:
        docs = []

        for p in file_paths:
            if not p.lower().endswith(".csv"):
                continue

            with open(p, encoding="utf-8", newline="") as f:
                reader = csv.DictReader(f)

                #turn each row into key=value pairs
                for i, row in enumerate(reader):
                    text = ", ".join(f"{k}={v}" for k, v in row.items())
                    docs.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": Path(p).name,
                                "file_path": p,
                                "row": i,
                                "type": "csv"
                            }
                        )
                    )
        return docs

class TXTLoader(DocumentLoader):
    #simple text file loader, reads the whole file as one document
    
    def load(self, file_paths: List[str]) -> List[Document]:
        docs = []

        for p in file_paths:
            if not p.lower().endswith(".txt"):
                continue

            with open(p, encoding="utf-8") as f:
                text = f.read().strip()

            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": Path(p).name,
                        "file_path": p,
                        "type": "txt"
                    }
                )
            )
        return docs

class XLSMLoader(DocumentLoader):
    #handles Excel files (.xlsm and .xlsx) - works similar to CSV loader
    
    def load(self, file_paths: List[str]) -> List[Document]:
        docs = []

        try:
            import openpyxl
        except ImportError:
            raise ImportError("You need openpyxl installed: pip install openpyxl")

        for p in file_paths:
            if not (p.lower().endswith(".xlsm") or p.lower().endswith(".xlsx")):
                continue

            wb = openpyxl.load_workbook(p, data_only=True)
            sheet = wb.active

            #grab headers from first row
            headers = []
            for cell in sheet[1]:
                headers.append(cell.value if cell.value is not None else "")

            #process each data row
            for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
                values = []
                for h, v in zip(headers, row):
                    values.append(f"{h}={v}")

                docs.append(
                    Document(
                        page_content=", ".join(values),
                        metadata={
                            "source": Path(p).name,
                            "file_path": p,
                            "row": i,
                            "type": "xlsm"
                        }
                    )
                )
        return docs